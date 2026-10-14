import json
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import ShopSettings
from apps.products.models import Product

from . import services
from .models import Sale, SaleItem, StockMovement


def configure_shop(**overrides):
    shop = ShopSettings.get_solo()
    shop.shop_name = "Test Shop"
    shop.is_configured = True
    for key, value in overrides.items():
        setattr(shop, key, value)
    shop.save()
    return shop


class TaxCalculationTests(TestCase):
    def setUp(self):
        configure_shop()
        self.product = Product.objects.create(
            name="Widget", selling_price="100.00", gst_rate="18", stock_quantity=100
        )

    def _line(self, mode, price, gst, qty=1):
        return services.calculate_line(qty, price, gst, mode)

    def test_exclusive(self):
        line = self._line(ShopSettings.PRICING_EXCLUSIVE, "100.00", "18")
        self.assertEqual(line.taxable_amount, Decimal("100.00"))
        self.assertEqual(line.tax_amount, Decimal("18.00"))
        self.assertEqual(line.line_total, Decimal("118.00"))

    def test_inclusive(self):
        line = self._line(ShopSettings.PRICING_INCLUSIVE, "118.00", "18")
        self.assertEqual(line.taxable_amount, Decimal("100.00"))
        self.assertEqual(line.tax_amount, Decimal("18.00"))
        self.assertEqual(line.line_total, Decimal("118.00"))

    def test_zero_gst(self):
        line = self._line(ShopSettings.PRICING_EXCLUSIVE, "50.00", "0")
        self.assertEqual(line.tax_amount, Decimal("0.00"))
        self.assertEqual(line.line_total, Decimal("50.00"))

    def test_rounding_exclusive(self):
        line = self._line(ShopSettings.PRICING_EXCLUSIVE, "99.99", "18")
        self.assertEqual(line.tax_amount, Decimal("18.00"))
        self.assertEqual(line.line_total, Decimal("117.99"))

    def test_cart_totals_sum_rounded_lines(self):
        p1 = Product.objects.create(
            name="A", selling_price="99.99", gst_rate="18", stock_quantity=10
        )
        p2 = Product.objects.create(
            name="B", selling_price="0.05", gst_rate="18", stock_quantity=10
        )
        cart = services.calculate_cart(
            [{"product": p1, "quantity": 1}, {"product": p2, "quantity": 1}],
            ShopSettings.PRICING_EXCLUSIVE,
        )
        self.assertEqual(cart.grand_total, Decimal("118.05"))


class CheckoutTests(TestCase):
    def setUp(self):
        configure_shop(pricing_mode=ShopSettings.PRICING_EXCLUSIVE)
        self.product = Product.objects.create(
            name="Rice", barcode="111", selling_price="100.00", gst_rate="5", stock_quantity=10
        )

    def _checkout(self, items, token="tok-1"):
        return services.perform_checkout(cart_payload=items, checkout_token=token)

    def test_successful_sale_deducts_stock_once(self):
        sale, created = self._checkout(
            [{"product_id": self.product.pk, "quantity": 3}]
        )
        self.assertTrue(created)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 7)
        self.assertEqual(Sale.objects.count(), 1)
        self.assertEqual(SaleItem.objects.count(), 1)
        self.assertEqual(
            StockMovement.objects.filter(reason=StockMovement.REASON_SALE).count(), 1
        )
        self.assertEqual(sale.grand_total, Decimal("315.00"))

    def test_duplicate_token_returns_original_sale(self):
        first, _ = self._checkout([{"product_id": self.product.pk, "quantity": 2}])
        second, created = self._checkout(
            [{"product_id": self.product.pk, "quantity": 2}], token="tok-1"
        )
        self.assertFalse(created)
        self.assertEqual(first.pk, second.pk)
        self.assertEqual(Sale.objects.count(), 1)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 8)

    def test_insufficient_stock_blocks_checkout_and_leaves_nothing(self):
        with self.assertRaises(services.CheckoutError):
            self._checkout([{"product_id": self.product.pk, "quantity": 50}])
        self.assertEqual(Sale.objects.count(), 0)
        self.assertEqual(SaleItem.objects.count(), 0)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 10)

    def test_inactive_product_cannot_be_sold(self):
        self.product.is_active = False
        self.product.save()
        with self.assertRaises(services.CheckoutError):
            self._checkout([{"product_id": self.product.pk, "quantity": 1}])
        self.assertEqual(Sale.objects.count(), 0)

    def test_partial_failure_rolls_back_all_lines(self):
        second = Product.objects.create(
            name="Oil", selling_price="200.00", gst_rate="5", stock_quantity=1
        )
        with self.assertRaises(services.CheckoutError):
            self._checkout(
                [
                    {"product_id": self.product.pk, "quantity": 2},
                    {"product_id": second.pk, "quantity": 5},
                ]
            )
        self.product.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 10)
        self.assertEqual(second.stock_quantity, 1)
        self.assertEqual(StockMovement.objects.count(), 0)

    def test_server_ignores_client_supplied_prices(self):
        # Prices come from the product record, never from the payload.
        sale, _ = self._checkout(
            [{"product_id": self.product.pk, "quantity": 1}]
        )
        item = sale.items.first()
        self.assertEqual(item.unit_price, Decimal("100.00"))

    def test_snapshot_survives_product_edit(self):
        sale, _ = self._checkout([{"product_id": self.product.pk, "quantity": 1}])
        self.product.name = "Basmati Rice"
        self.product.selling_price = Decimal("999.00")
        self.product.save()
        item = sale.items.first()
        self.assertEqual(item.product_name_snapshot, "Rice")
        self.assertEqual(item.unit_price, Decimal("100.00"))

    def test_deactivating_product_keeps_historical_sale(self):
        sale, _ = self._checkout([{"product_id": self.product.pk, "quantity": 1}])
        self.product.is_active = False
        self.product.save()
        self.assertTrue(Sale.objects.filter(pk=sale.pk).exists())

    def test_checkout_records_payment_method(self):
        sale, _ = services.perform_checkout(
            cart_payload=[{"product_id": self.product.pk, "quantity": 1}],
            checkout_token="pay-upi",
            payment_method=Sale.PAYMENT_UPI,
        )
        self.assertEqual(sale.payment_method, Sale.PAYMENT_UPI)

    def test_checkout_rejects_unknown_payment_method(self):
        with self.assertRaises(services.CheckoutError):
            services.perform_checkout(
                cart_payload=[{"product_id": self.product.pk, "quantity": 1}],
                checkout_token="pay-bogus",
                payment_method="bitcoin",
            )
        self.assertEqual(Sale.objects.count(), 0)


class CheckoutViewTests(TestCase):
    def setUp(self):
        configure_shop(pricing_mode=ShopSettings.PRICING_EXCLUSIVE)
        self.product = Product.objects.create(
            name="Tea", barcode="500", selling_price="50.00", gst_rate="5", stock_quantity=20
        )

    def _post(self, payload):
        return self.client.post(
            reverse("billing:checkout"),
            data=json.dumps(payload),
            content_type="application/json",
        )

    def test_checkout_creates_sale(self):
        response = self._post(
            {
                "token": "view-token",
                "items": [{"product_id": self.product.pk, "quantity": 2}],
            }
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["ok"])
        self.assertEqual(Sale.objects.count(), 1)

    def test_checkout_rejects_empty_cart(self):
        response = self._post({"token": "t", "items": []})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(response.json()["ok"])

    def test_checkout_rejects_overselling(self):
        response = self._post(
            {
                "token": "t2",
                "items": [{"product_id": self.product.pk, "quantity": 999}],
            }
        )
        self.assertEqual(response.status_code, 400)

    def test_checkout_requires_token(self):
        response = self._post(
            {"items": [{"product_id": self.product.pk, "quantity": 1}]}
        )
        self.assertEqual(response.status_code, 400)


class POSAndSalesViewTests(TestCase):
    def setUp(self):
        configure_shop()
        self.product = Product.objects.create(
            name="Sugar", barcode="900", selling_price="40.00", gst_rate="5", stock_quantity=30
        )

    def test_pos_page_renders(self):
        response = self.client.get(reverse("billing:pos"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Current Bill")
        self.assertContains(response, "Complete Sale")
        self.assertContains(response, "<svg")
        self.assertNotContains(response, "&lt;svg")

    def test_search_endpoint(self):
        response = self.client.get(reverse("billing:product_search"), {"q": "900"})
        data = response.json()
        self.assertEqual(len(data["results"]), 1)
        self.assertEqual(data["exact_id"], self.product.pk)

    def test_barcode_lookup_unknown(self):
        response = self.client.get(reverse("billing:barcode_lookup"), {"barcode": "nope"})
        self.assertFalse(response.json()["found"])

    def test_sale_list_and_detail(self):
        services.perform_checkout(
            cart_payload=[{"product_id": self.product.pk, "quantity": 1}],
            checkout_token="sale-view",
        )
        sale = Sale.objects.get()
        listing = self.client.get(reverse("billing:sale_list"))
        self.assertContains(listing, sale.bill_number)
        detail = self.client.get(
            reverse("billing:sale_detail", args=[sale.bill_number])
        )
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, sale.bill_number)

    def test_reprint_does_not_change_stock(self):
        services.perform_checkout(
            cart_payload=[{"product_id": self.product.pk, "quantity": 2}],
            checkout_token="reprint",
        )
        self.product.refresh_from_db()
        stock_after_sale = self.product.stock_quantity
        sale = Sale.objects.get()
        self.client.get(reverse("billing:sale_detail", args=[sale.bill_number]))
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, stock_after_sale)
        self.assertEqual(Sale.objects.count(), 1)

    def _sale_with_payment(self, token, payment):
        services.perform_checkout(
            cart_payload=[{"product_id": self.product.pk, "quantity": 1}],
            checkout_token=token,
            payment_method=payment,
        )
        return Sale.objects.get(checkout_token=token)

    def test_sale_list_filters_by_payment_method(self):
        cash_sale = self._sale_with_payment("cash-sale", Sale.PAYMENT_CASH)
        upi_sale = self._sale_with_payment("upi-sale", Sale.PAYMENT_UPI)

        listing = self.client.get(reverse("billing:sale_list"), {"payment": "upi"})
        self.assertContains(listing, upi_sale.bill_number)
        self.assertNotContains(listing, cash_sale.bill_number)

        listing = self.client.get(reverse("billing:sale_list"), {"payment": "cash"})
        self.assertContains(listing, cash_sale.bill_number)
        self.assertNotContains(listing, upi_sale.bill_number)

    def test_sale_list_ignores_unknown_payment_filter(self):
        sale = self._sale_with_payment("unknown-filter", Sale.PAYMENT_CASH)
        listing = self.client.get(
            reverse("billing:sale_list"), {"payment": "bitcoin"}
        )
        self.assertEqual(listing.status_code, 200)
        self.assertContains(listing, sale.bill_number)

    def test_print_mode_triggers_autoprint(self):
        services.perform_checkout(
            cart_payload=[{"product_id": self.product.pk, "quantity": 1}],
            checkout_token="print-mode",
        )
        sale = Sale.objects.get()
        url = reverse("billing:sale_detail", args=[sale.bill_number])
        normal = self.client.get(url)
        self.assertNotContains(normal, 'window.addEventListener("load"')
        printing = self.client.get(url, {"print": "1"})
        self.assertEqual(printing.status_code, 200)
        self.assertContains(printing, 'window.addEventListener("load"')
        self.assertContains(printing, sale.bill_number)


class SaleExportTests(TestCase):
    def setUp(self):
        configure_shop(pricing_mode=ShopSettings.PRICING_EXCLUSIVE)
        self.product = Product.objects.create(
            name="Pen", barcode="700", selling_price="10.00", gst_rate="0",
            stock_quantity=100,
        )

    def _sale(self, token, payment):
        sale, _ = services.perform_checkout(
            cart_payload=[{"product_id": self.product.pk, "quantity": 1}],
            checkout_token=token,
            payment_method=payment,
        )
        return sale

    def _export(self, **params):
        return self.client.get(reverse("billing:sale_export"), params)

    def _sheet_text(self, response):
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        return "\n".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )

    def test_export_returns_xlsx_attachment(self):
        self._sale("export-1", Sale.PAYMENT_CASH)
        response = self._export()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook.active.title, "Bill history")

    def test_export_respects_payment_filter(self):
        cash_sale = self._sale("export-cash", Sale.PAYMENT_CASH)
        upi_sale = self._sale("export-upi", Sale.PAYMENT_UPI)

        text = self._sheet_text(self._export(payment="upi"))
        self.assertIn(upi_sale.bill_number, text)
        self.assertNotIn(cash_sale.bill_number, text)

        text = self._sheet_text(self._export(payment="cash"))
        self.assertIn(cash_sale.bill_number, text)
        self.assertNotIn(upi_sale.bill_number, text)

    def test_export_respects_date_range(self):
        sale = self._sale("export-dated", Sale.PAYMENT_CASH)
        day = timezone.localtime(sale.created_at).strftime("%Y-%m-%d")

        inside = self._sheet_text(self._export(date_from=day, date_to=day))
        self.assertIn(sale.bill_number, inside)

        outside = self._sheet_text(
            self._export(date_from="2000-01-01", date_to="2000-01-02")
        )
        self.assertNotIn(sale.bill_number, outside)

    def test_export_includes_totals_row(self):
        self._sale("export-total", Sale.PAYMENT_CASH)
        text = self._sheet_text(self._export())
        self.assertIn("TOTAL", text)

    def test_export_rejects_post(self):
        response = self.client.post(reverse("billing:sale_export"))
        self.assertEqual(response.status_code, 405)
