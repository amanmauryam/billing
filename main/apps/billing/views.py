"""POS, checkout, and sales history views."""

import json
from datetime import datetime, time
from decimal import Decimal
from urllib.parse import urlencode

from django.contrib import messages
from django.core.paginator import Paginator
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_GET, require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.products.models import Product

from . import services
from .models import Sale

DATE_FORMAT = "%Y-%m-%d"
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _product_payload(product):
    return {
        "id": product.pk,
        "name": product.name,
        "barcode": product.barcode or "",
        "price": str(product.selling_price),
        "gst_rate": str(product.gst_rate),
        "stock": product.stock_quantity,
        "is_active": product.is_active,
        "stock_status": product.stock_status,
    }


@ensure_csrf_cookie
def pos_view(request):
    from apps.core.models import ShopSettings

    shop = ShopSettings.get_solo()
    return render(
        request,
        "pos/pos.html",
        {
            "page_title": "New Bill",
            "pricing_mode": shop.pricing_mode,
            "pricing_inclusive": shop.pricing_mode == ShopSettings.PRICING_INCLUSIVE,
        },
    )


@require_GET
def product_search(request):
    query = (request.GET.get("q") or "").strip()
    if not query:
        return JsonResponse({"results": []})

    exact = Product.objects.active().filter(barcode__iexact=query).first()
    results = list(Product.objects.active().search(query)[:20])
    if exact is not None and all(p.pk != exact.pk for p in results):
        results.insert(0, exact)

    return JsonResponse(
        {
            "results": [_product_payload(p) for p in results],
            "exact_id": exact.pk if exact is not None else None,
        }
    )


@require_GET
def barcode_lookup(request):
    barcode = (request.GET.get("barcode") or "").strip()
    if not barcode:
        return JsonResponse({"found": False, "error": "Empty barcode."}, status=400)
    product = Product.objects.filter(barcode__iexact=barcode).first()
    if product is None:
        return JsonResponse({"found": False, "barcode": barcode})
    if not product.is_active:
        return JsonResponse(
            {
                "found": False,
                "error": f"'{product.name}' is inactive and cannot be sold.",
                "barcode": barcode,
            }
        )
    return JsonResponse({"found": True, "product": _product_payload(product)})


@require_POST
def checkout(request):
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except (json.JSONDecodeError, UnicodeDecodeError):
        return JsonResponse(
            {"ok": False, "error": "Malformed request."}, status=400
        )

    token = payload.get("token")
    items = payload.get("items")
    payment_method = payload.get("payment_method") or ""

    try:
        sale, created = services.perform_checkout(
            cart_payload=items,
            checkout_token=token,
            payment_method=payment_method,
        )
    except services.CheckoutError as exc:
        return JsonResponse({"ok": False, "error": exc.message}, status=400)

    return JsonResponse(
        {
            "ok": True,
            "created": created,
            "bill_number": sale.bill_number,
            "grand_total": str(sale.grand_total),
            "receipt_url": reverse("billing:sale_detail", args=[sale.bill_number]),
            "new_bill_url": reverse("billing:pos"),
        }
    )


def _clean(value):
    return (value or "").strip()


def _parse_day(value, *, end_of_day):
    """Parse a ``YYYY-MM-DD`` string into an aware boundary datetime."""
    try:
        day = datetime.strptime(value, DATE_FORMAT).date()
    except (TypeError, ValueError):
        return None
    bound = time.max if end_of_day else time.min
    return timezone.make_aware(datetime.combine(day, bound))


def _apply_sale_filters(request):
    """Return ``(queryset, filters, warnings)`` for the bill history views.

    Both the HTML list and the Excel export use this so the report always
    matches exactly what the shopkeeper sees on screen.
    """
    filters = {
        "bill_number": _clean(request.GET.get("bill")),
        "payment": _clean(request.GET.get("payment")),
        "date_from": _clean(request.GET.get("date_from")),
        "date_to": _clean(request.GET.get("date_to")),
    }
    # Backwards compatibility with the original single-day ``date`` filter.
    legacy_date = _clean(request.GET.get("date"))
    if legacy_date:
        filters["date_from"] = filters["date_from"] or legacy_date
        filters["date_to"] = filters["date_to"] or legacy_date

    warnings = []
    sales = Sale.objects.completed()

    if filters["bill_number"]:
        sales = sales.filter(bill_number__icontains=filters["bill_number"])

    if filters["payment"]:
        if filters["payment"] in dict(Sale.PAYMENT_CHOICES):
            sales = sales.filter(payment_method=filters["payment"])
        else:
            warnings.append("Ignoring an unknown payment filter.")
            filters["payment"] = ""

    start = _parse_day(filters["date_from"], end_of_day=False)
    if filters["date_from"] and start is None:
        warnings.append("Ignoring an invalid start date.")
        filters["date_from"] = ""
    end = _parse_day(filters["date_to"], end_of_day=True)
    if filters["date_to"] and end is None:
        warnings.append("Ignoring an invalid end date.")
        filters["date_to"] = ""

    if start is not None and end is not None and start > end:
        warnings.append("The date range was reversed; start and end were swapped.")
        start, end = end, start

    if start is not None:
        sales = sales.filter(created_at__gte=start)
    if end is not None:
        sales = sales.filter(created_at__lte=end)

    return sales, filters, warnings


def _export_query_string(filters):
    params = {}
    if filters["bill_number"]:
        params["bill"] = filters["bill_number"]
    if filters["payment"]:
        params["payment"] = filters["payment"]
    if filters["date_from"]:
        params["date_from"] = filters["date_from"]
    if filters["date_to"]:
        params["date_to"] = filters["date_to"]
    return urlencode(params)


def _filter_summary(filters):
    parts = []
    if filters["bill_number"]:
        parts.append(f"Bill contains '{filters['bill_number']}'")
    if filters["payment"]:
        label = dict(Sale.PAYMENT_CHOICES).get(filters["payment"], filters["payment"])
        parts.append(f"Payment: {label}")
    if (
        filters["date_from"]
        and filters["date_to"]
        and filters["date_from"] == filters["date_to"]
    ):
        parts.append(f"Date: {filters['date_from']}")
    else:
        if filters["date_from"]:
            parts.append(f"From: {filters['date_from']}")
        if filters["date_to"]:
            parts.append(f"To: {filters['date_to']}")
    return " · ".join(parts) if parts else "All transactions"


def sale_list(request):
    sales, filters, warnings = _apply_sale_filters(request)
    for warning in warnings:
        messages.warning(request, warning)

    paginator = Paginator(sales, 20)
    page = paginator.get_page(request.GET.get("page"))

    return render(
        request,
        "sales/list.html",
        {
            "page_title": "Sales",
            "page": page,
            "payment_choices": Sale.PAYMENT_CHOICES,
            "total_count": paginator.count,
            "has_sales": Sale.objects.completed().exists(),
            "has_filters": any(filters.values()),
            "export_query": _export_query_string(filters),
            **filters,
        },
    )


@require_GET
def sale_export(request):
    """Export the currently filtered bill history as an ``.xlsx`` report."""
    from apps.core.models import ShopSettings

    sales, filters, _warnings = _apply_sale_filters(request)
    shop = ShopSettings.get_solo()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bill history"

    title_font = Font(bold=True, size=14)
    meta_font = Font(color="697386", size=10)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="315EFB")
    bold = Font(bold=True)
    money_format = "#,##0.00"
    center = Alignment(horizontal="center")
    left = Alignment(horizontal="left")
    right = Alignment(horizontal="right")

    sheet["A1"] = shop.shop_name or "Shop"
    sheet["A1"].font = title_font
    sheet["A2"] = "Bill history report"
    sheet["A2"].font = Font(bold=True, size=11)
    sheet["A3"] = _filter_summary(filters)
    sheet["A3"].font = meta_font
    sheet["A4"] = f"Generated {timezone.localtime():%d %b %Y, %H:%M}"
    sheet["A4"].font = meta_font

    headers = [
        "Bill number",
        "Date & time",
        "Payment mode",
        "Items",
        "Subtotal",
        "Taxable",
        "GST",
        "Grand total",
    ]
    header_row = 6
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    totals = {5: Decimal("0.00"), 6: Decimal("0.00"), 7: Decimal("0.00"), 8: Decimal("0.00")}
    total_items = 0
    row = header_row + 1

    for sale in sales.iterator():
        sheet.cell(row=row, column=1, value=sale.bill_number)
        sheet.cell(
            row=row,
            column=2,
            value=timezone.localtime(sale.created_at).strftime("%d %b %Y, %H:%M"),
        )
        sheet.cell(
            row=row,
            column=3,
            value=sale.get_payment_method_display() or "Not recorded",
        )
        sheet.cell(row=row, column=4, value=sale.item_count).alignment = center
        for column, amount in (
            (5, sale.subtotal),
            (6, sale.taxable_total),
            (7, sale.tax_total),
            (8, sale.grand_total),
        ):
            cell = sheet.cell(row=row, column=column, value=float(amount))
            cell.number_format = money_format
        total_items += sale.item_count
        for column, amount in (
            (5, sale.subtotal),
            (6, sale.taxable_total),
            (7, sale.tax_total),
            (8, sale.grand_total),
        ):
            totals[column] += amount
        row += 1

    if row > header_row + 1:
        sheet.cell(row=row, column=1, value="TOTAL").font = bold
        sheet.cell(row=row, column=4, value=total_items).alignment = center
        for column, amount in totals.items():
            cell = sheet.cell(row=row, column=column, value=float(amount))
            cell.number_format = money_format
            cell.font = bold

    widths = [22, 22, 16, 8, 14, 14, 12, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    for column in range(1, len(headers) + 1):
        sheet.cell(row=header_row, column=column).alignment = center
    for anchored_row in (1, 2, 3, 4):
        sheet.cell(row=anchored_row, column=1).alignment = left
    for column in range(5, 9):
        for data_row in range(header_row + 1, row + 1):
            sheet.cell(row=data_row, column=column).alignment = right

    filename = f"bill-history-{timezone.localtime():%Y%m%d-%H%M}.xlsx"
    response = HttpResponse(content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def sale_detail(request, bill_number):
    sale = get_object_or_404(
        Sale.objects.prefetch_related("items"), bill_number=bill_number
    )
    return render(
        request,
        "sales/detail.html",
        {
            "page_title": f"Bill {sale.bill_number}",
            "sale": sale,
            "items": sale.items.all(),
            "print_mode": request.GET.get("print") == "1",
        },
    )
